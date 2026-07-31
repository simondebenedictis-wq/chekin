#!/usr/bin/env python3
"""
Weekly data updater for dashboard/index.html.

Usage:
    python3 scripts/update_dashboard_data.py <path-to-sheet-export.xlsx>

Takes an .xlsx export of the "Chekin Dashboard Data (Public)" Google Sheet
(File > Download > Microsoft Excel) and regenerates the STATIC_ROWS block
baked into dashboard/index.html, plus updates the "as of <date>" labels.

Column layout notes (verified against the sheet as of Jul 2026):
  - Each sheet's real column positions are configured in SHEETS below.
    Layouts have drifted between sheet copies before (REMOTE ACCESS lost a
    filler block, SECURITY DEPOSIT gained one) — if this script's sanity
    checks warn about a mismatch, open the .xlsx and re-check the header
    rows for that tab before trusting the output.
  - Header rows: 3 for BOOKABLE UNITS (data starts row 4), 2 for every
    other tab (data starts row 3).
"""
import sys
import re
import json
import datetime
import openpyxl

SHEETS = {
    # key: (sheet name, number of columns to read, data start row)
    'bu':       ('BOOKABLE UNITS', 41, 4),
    'idv':      ('IDENTITY VERFICATION', 43, 3),
    'remote':   ('REMOTE ACCESS', 35, 3),
    'branded':  ('BRANDED GUEST APP', 37, 3),
    'guide':    ('DIGITAL GUIDEBOOK', 35, 3),
    'einvoice': ('E-INVOICING', 32, 3),
    'tax':      ('TOURIST TAXES ', 41, 3),
    'deposit':  ('SECURITY DEPOSIT', 59, 3),
    'damage':   ('DAMAGE PROTECTION', 35, 3),
    'liveness': ('LIVENESS', 29, 3),
    'inbox':    ('SMART INBOX', 31, 3),
}

# Base block column indices (0-based, col A = index 0) used ONLY for the
# sanity check below — must mirror the OFFSETS/parse* logic in
# dashboard/index.html. SMART INBOX has a fully different layout (no
# tot/H/VR/UNK block at these positions), so it gets its own entry.
BASE_OFFSETS = {
    'bu':       {'tot': 4, 'H': 6, 'VR': 7, 'UNK': 8},
    'default':  {'tot': 5, 'H': 7, 'VR': 8, 'UNK': 9},
    'inbox':    {'tot': 4, 'H': 9, 'VR': 10, 'UNK': 11},  # tot=TOTAL USERS, not H+VR+UNK-summable the same way
}

NWEEKS = 30  # how many weeks of history to bake in for trend charts


def dump_sheet(ws, ncols, start_row, nweeks):
    rows = []
    for r in range(start_row, start_row + nweeks):
        cells = [ws.cell(row=r, column=c).value for c in range(1, ncols + 1)]
        if cells[0] is None or not isinstance(cells[0], datetime.datetime):
            break
        rows.append(cells)
    return rows


def to_jsval(v):
    if v is None:
        return 'null'
    if isinstance(v, datetime.datetime):
        return f'[{v.year},{v.month - 1},{v.day}]'
    if isinstance(v, (int, float)):
        if v != v:  # NaN
            return 'null'
        return repr(v)
    if isinstance(v, str):
        return json.dumps(v)
    return 'null'


def num(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return v if v == v else 0
    if isinstance(v, str):
        s = v.strip()
        if not s or s == '-' or s.startswith('#'):
            return 0
        try:
            return float(s.replace(',', ''))
        except ValueError:
            return 0
    return 0


def sanity_check(key, rows, off):
    """Cheap arithmetic checks to catch a column-layout drift early.
    Skipped for 'inbox': its H/VR/UNK columns are a sub-breakdown that
    doesn't sum to its "tot", so the tot ≈ H+VR+UNK heuristic doesn't apply.
    """
    if not rows:
        print(f'  [{key}] WARNING: no data rows found — check sheet name / header row count')
        return
    latest = rows[0]
    tot = num(latest[off['tot']])
    h = num(latest[off['H']])
    vr = num(latest[off['VR']])
    unk = num(latest[off['UNK']])
    if key == 'inbox':
        print(f'  [{key}] latest={latest[0].date()} totalUsers={tot:.0f} (H/VR/UNK breakdown not checked)')
        return
    diff = tot - (h + vr + unk)
    flag = '' if abs(diff) <= max(5, tot * 0.02) else '  <-- CHECK OFFSETS (H+VR+UNK far from TOT)'
    print(f'  [{key}] latest={latest[0].date()} tot={tot:.0f} H={h:.0f} VR={vr:.0f} UNK={unk:.0f}{flag}')


def build_static_rows(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    out_lines = ['const STATIC_ROWS = {']
    latest_date = None
    print('Extracted (latest week per sheet):')
    for key, (sheet_name, ncols, start_row) in SHEETS.items():
        if sheet_name not in wb.sheetnames:
            print(f'  [{key}] WARNING: sheet "{sheet_name}" not found in workbook — skipping')
            out_lines.append(f'  {key}: [],')
            continue
        ws = wb[sheet_name]
        rows = dump_sheet(ws, ncols, start_row, NWEEKS)
        off = BASE_OFFSETS.get(key, BASE_OFFSETS['default'])
        sanity_check(key, rows, off)
        if rows and (latest_date is None or rows[0][0] > latest_date):
            latest_date = rows[0][0]
        out_lines.append(f'  {key}: [')
        for cells in rows:
            out_lines.append('    [' + ','.join(to_jsval(v) for v in cells) + '],')
        out_lines.append('  ],')
    out_lines.append('};')
    return '\n'.join(out_lines), latest_date


def splice_into_html(html_path, static_js, latest_date):
    with open(html_path) as f:
        html = f.read()

    start_marker = 'const STATIC_ROWS = {'
    end_marker = '\nfunction hydrateFromStatic()'
    start_idx = html.index(start_marker)
    end_idx = html.index(end_marker, start_idx)
    html = html[:start_idx] + static_js + html[end_idx:]

    if latest_date:
        date_str = latest_date.strftime('%b %-d, %Y') if sys.platform != 'win32' else latest_date.strftime('%b %#d, %Y')
        # Update the header default text and footer snapshot label.
        html = re.sub(r'Week of [A-Za-z]{3} \d{1,2}, \d{4}', f'Week of {date_str}', html)
        html = re.sub(r'Data snapshot: [A-Za-z]{3} \d{1,2}, \d{4}', f'Data snapshot: {date_str}', html)

    with open(html_path, 'w') as f:
        f.write(html)
    return date_str if latest_date else None


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    xlsx_path = sys.argv[1]
    html_path = 'dashboard/index.html'

    static_js, latest_date = build_static_rows(xlsx_path)
    date_str = splice_into_html(html_path, static_js, latest_date)

    print()
    print(f'Updated {html_path} with data through {date_str or "(unknown — no date found)"}.')
    print('Next: review the sanity-check output above, then commit + republish the artifact.')


if __name__ == '__main__':
    main()
